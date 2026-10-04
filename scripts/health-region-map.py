"""Build the premium-region SVG map from official municipality geometries.

Generation-only dependencies: openpyxl, pyshp, shapely (>=2.1). No frontend or
API runtime dependency. Install in an isolated Python environment, then run:
  python scripts/health-region-map.py
An existing dependency directory can be supplied with --python-dependencies.
The immutable boundary archive is checksum-verified and cached in the OS temp
directory. Commune mergers are read from the supplied official Priminfo XLSX.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import pathlib
import re
import sys
import tempfile
import urllib.request
import zipfile
from collections import defaultdict


BOUNDARY_URL = (
    "https://data.geo.admin.ch/ch.swisstopo.swissboundaries3d/"
    "swissboundaries3d_2026-01/swissboundaries3d_2026-01_2056_5728.shp.zip"
)
BOUNDARY_SHA256 = "a467d7eec8395d746253c68095c75a2540b52064bc83a6b35cf101efb625733d"
REGIONS_URL = "https://www.priminfo.admin.ch/downloads/praemienregionen.xlsx"
ROOT = pathlib.Path(__file__).resolve().parents[1]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--python-dependencies", type=pathlib.Path)
    parser.add_argument(
        "--regions", type=pathlib.Path,
        default=ROOT / "data/health-insurance/2027/regions.xlsx",
    )
    parser.add_argument(
        "--boundaries", type=pathlib.Path,
        default=pathlib.Path(tempfile.gettempdir()) / "coverview-swissboundaries3d-2026.zip",
    )
    parser.add_argument(
        "--output", type=pathlib.Path,
        default=ROOT / "data/health-insurance/2027/regions-map.json",
    )
    args = parser.parse_args()
    if args.python_dependencies:
        sys.path.insert(0, str(args.python_dependencies.resolve()))

    import openpyxl
    import shapefile
    from shapely import coverage_is_valid, coverage_simplify, force_2d, make_valid
    from shapely.geometry import shape
    from shapely.ops import unary_union

    workbook_bytes = args.regions.read_bytes()
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    if "2027" not in str(workbook["Informations"]["A1"].value):
        raise ValueError("The region workbook must explicitly describe premium year 2027")
    communes: dict[int, tuple[str, int]] = {}
    for row in workbook["A_COM"].iter_rows(min_row=6, values_only=True):
        if not isinstance(row[0], (int, float)):
            continue
        bfs = int(row[0])
        region = (str(row[1]), int(row[3]))
        if bfs in communes and communes[bfs] != region:
            raise ValueError(f"Conflicting regions for municipality {bfs}")
        communes[bfs] = region

    mutations: dict[int, int] = {}
    for row in workbook["COM-MUT"].iter_rows(values_only=True):
        if not isinstance(row[0], (int, float)) or int(row[0]) > 2027:
            continue
        match = re.match(r"fus\s+(\d+)\b", str(row[4]))
        if match and int(row[1]) != int(match[1]):
            mutations[int(row[1])] = int(match[1])
    workbook.close()

    if not args.boundaries.exists():
        with urllib.request.urlopen(BOUNDARY_URL, timeout=60) as response:
            payload = response.read(100_000_001)
        if len(payload) > 100_000_000:
            raise ValueError("Boundary archive exceeds 100 MB limit")
        args.boundaries.parent.mkdir(parents=True, exist_ok=True)
        args.boundaries.write_bytes(payload)
    archive_bytes = args.boundaries.read_bytes()
    digest = hashlib.sha256(archive_bytes).hexdigest()
    if digest != BOUNDARY_SHA256:
        raise ValueError(f"Unexpected boundary archive checksum: {digest}")

    groups = defaultdict(list)
    group_communes = defaultdict(set)
    source_communes = set()
    mapped_communes = set()
    applied_mutations = {}
    ignored = defaultdict(int)
    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        base = "swissBOUNDARIES3D_1_5_TLM_HOHEITSGEBIET"
        with shapefile.Reader(
            shp=io.BytesIO(archive.read(base + ".shp")),
            dbf=io.BytesIO(archive.read(base + ".dbf")),
            shx=io.BytesIO(archive.read(base + ".shx")),
            encoding="utf-8",
        ) as reader:
            for feature in reader.iterShapeRecords():
                attrs = feature.record.as_dict()
                if attrs["ICC"] != "CH" or attrs["OBJEKTART"] != "Gemeindegebiet":
                    ignored[f"{attrs['ICC']}:{attrs['OBJEKTART']}"] += 1
                    continue
                old_bfs = int(attrs["BFS_NUMMER"])
                bfs = mutations.get(old_bfs, old_bfs)
                if bfs not in communes:
                    raise ValueError(f"Unmapped municipality {old_bfs} -> {bfs}: {attrs['NAME']}")
                if old_bfs != bfs:
                    applied_mutations[old_bfs] = bfs
                source_communes.add(old_bfs)
                mapped_communes.add(bfs)
                key = communes[bfs]
                geom = force_2d(shape(feature.shape.__geo_interface__))
                if not geom.is_valid:
                    geom = make_valid(geom)
                groups[key].append(geom)
                group_communes[key].add(bfs)

    missing = set(communes) - mapped_communes
    if missing:
        raise ValueError(f"Region workbook municipalities missing from map: {sorted(missing)}")
    if len(groups) != 42:
        raise ValueError(f"Expected 42 Swiss premium regions, got {len(groups)}")

    keys = sorted(groups)
    dissolved = [unary_union(groups[key]) for key in keys]
    if not all(geom.is_valid and not geom.is_empty for geom in dissolved):
        raise ValueError("Invalid dissolved premium-region geometry")
    if not coverage_is_valid(dissolved):
        raise ValueError("Region geometries overlap or have inconsistent shared boundaries")
    # Shared borders are simplified together, avoiding cracks between regions.
    simplified = list(coverage_simplify(dissolved, 150.0))
    if not coverage_is_valid(simplified):
        raise ValueError("Simplification damaged shared region boundaries")

    country = unary_union(dissolved)
    minx, miny, maxx, maxy = country.bounds
    width, height, padding = 1000, 650, 24
    scale = min((width - padding * 2) / (maxx - minx), (height - padding * 2) / (maxy - miny))
    offsetx = (width - (maxx - minx) * scale) / 2
    offsety = (height - (maxy - miny) * scale) / 2

    def point(x: float, y: float) -> list[float]:
        return [round(offsetx + (x - minx) * scale, 2), round(offsety + (maxy - y) * scale, 2)]

    def path(geom) -> str:
        polygons = [geom] if geom.geom_type == "Polygon" else list(geom.geoms)
        chunks = []
        for polygon in polygons:
            if polygon.geom_type != "Polygon":
                raise ValueError(f"Unexpected geometry type: {polygon.geom_type}")
            for ring in [polygon.exterior, *polygon.interiors]:
                coords = [point(x, y) for x, y in ring.coords]
                chunks.append("M" + "L".join(f"{x:g},{y:g}" for x, y in coords[:-1]) + "Z")
        return "".join(chunks)

    regions = []
    for key, precise, geom in zip(keys, dissolved, simplified, strict=True):
        canton, region = key
        bminx, bminy, bmaxx, bmaxy = precise.bounds
        label = geom.representative_point()
        regions.append({
            "id": f"{canton}-{region}", "canton": canton, "region": region,
            "path": path(geom), "centroid": point(precise.centroid.x, precise.centroid.y),
            "labelPoint": point(label.x, label.y),
            "bounds": [*point(bminx, bmaxy), *point(bmaxx, bminy)],
            "municipalityCount": len(group_communes[key]),
        })

    output = {
        "year": 2027, "width": width, "height": height, "viewBox": "0 0 1000 650",
        "attribution": "© swisstopo · Régions de primes : OFSP / Priminfo 2027",
        "boundaryDate": "2026-01-01", "regionValidFrom": "2027-01-01",
        "regionValidTo": "2027-12-31", "fillRule": "evenodd",
        "method": "Union des communes OFS par canton/région après application des fusions officielles COM-MUT ; simplification topologique commune à 150 m, projection LV95.",
        "sources": [
            {"name": "swissBOUNDARIES3D 2026", "url": BOUNDARY_URL, "sha256": digest,
             "licenseUrl": "https://www.swisstopo.admin.ch/fr/conditions-utilisation-geodonnees-et-geoservices-gratuit"},
            {"name": "Régions de primes 2027, A_COM et COM-MUT", "url": REGIONS_URL,
             "sha256": hashlib.sha256(workbook_bytes).hexdigest()},
        ],
        "validation": {
            "sourceMunicipalities": len(source_communes),
            "targetMunicipalities": len(mapped_communes), "unmappedMunicipalities": [],
            "regionCount": len(regions), "cantonCount": len({key[0] for key in keys}),
            "sharedBoundariesValid": True,
            "appliedMunicipalityMergers": {str(k): v for k, v in sorted(applied_mutations.items())},
            "excludedNonMunicipalAndForeignFeatures": dict(sorted(ignored.items())),
        },
        "regions": regions,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "bytes": args.output.stat().st_size, **output["validation"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
